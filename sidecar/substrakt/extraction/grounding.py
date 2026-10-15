"""Ground extracted values in document evidence."""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from rapidfuzz import fuzz

from substrakt.index.words import page_text_with_index
from substrakt.model import Document, ImageRect, Locator, PdfRect, SheetRange, Word

THRESHOLD = 80.0



MAX_WINDOW = 8


@dataclass
class Evidence:

    locator: Locator
    words: list[Word] = field(default_factory=list)
    similarity: float = 0.0
    """Return how well the requested value matches the found text, from 0 to 100."""
    matched_text: str = ""

    @property
    def ocr_confidence(self) -> float | None:
        values = [w.confidence for w in self.words if w.confidence is not None]
        return min(values) if values else None


def _merge(words: list[Word]) -> Locator | None:
    if not words:
        return None

    first = words[0].locator
    if isinstance(first, SheetRange):
        
        from openpyxl.utils import get_column_letter

        from substrakt.utils.a1 import a1_to_row_column

        cells = [
            a1_to_row_column(w.locator.a1)
            for w in words
            if isinstance(w.locator, SheetRange)
        ]
        if not cells:
            return first
        z1, s1 = min(z for z, _ in cells), min(s for _, s in cells)
        z2, s2 = max(z for z, _ in cells), max(s for _, s in cells)
        span = (
            f"{get_column_letter(s1)}{z1}"
            if (z1, s1) == (z2, s2)
            else f"{get_column_letter(s1)}{z1}:{get_column_letter(s2)}{z2}"
        )
        return SheetRange(sheet=first.sheet, a1=span)

    rectangles = [w.locator for w in words if not isinstance(w.locator, SheetRange)]
    if not rectangles:
        return None

    
    result = rectangles[0]
    for other in rectangles[1:]:
        if other.page != result.page:
            continue  
        result = result.merged_with(other)
    return result


def _comparable(text: str) -> str:
    return re.sub(r"[^\w]", "", text).lower()





def from_position(
    doc: Document, page: int, start: int, end: int
) -> Evidence | None:
    _text, assignment = page_text_with_index(doc, page)
    if not assignment:
        return None

    involved = [
        word
        for w_start, w_ende, word in assignment
        if w_start < end and w_ende > start  
    ]
    if not involved:
        return None

    locator = _merge(involved)
    if locator is None:
        return None

    return Evidence(
        locator=locator,
        words=involved,
        similarity=100.0,  
        matched_text=" ".join(w.text for w in involved),
    )





def from_value(
    doc: Document, value: str, pages: list[int] | None = None
) -> Evidence | None:
    needle = _comparable(value)
    if not needle:
        return None

    best: Evidence | None = None
    pages_to_check = pages if pages is not None else [p.number for p in doc.pages]

    for page in pages_to_check:
        words = doc.words_on_page(page)
        if not words:
            continue

        
        
        for window_start in range(len(words)):
            for length in range(1, min(MAX_WINDOW, len(words) - window_start) + 1):
                window = words[window_start : window_start + length]
                candidate = _comparable("".join(w.text for w in window))
                if not candidate:
                    continue

                
                
                if len(candidate) > len(needle) * 2 + 10:
                    break

                similarity = fuzz.ratio(needle, candidate)
                if best is None or similarity > best.similarity:
                    locator = _merge(window)
                    if locator is not None:
                        best = Evidence(
                            locator=locator,
                            words=list(window),
                            similarity=similarity,
                            matched_text=" ".join(w.text for w in window),
                        )
                if similarity >= 100:
                    return best

    return best if best and best.similarity >= THRESHOLD else None


# --- Confidence ------------------------------------------------------------


def _in_region(evidence: Evidence, regions: list[Locator]) -> bool:
    if not regions:
        return False

    from substrakt.index.words import _rect_overlaps

    for region in regions:
        if isinstance(region, SheetRange) and isinstance(evidence.locator, SheetRange):
            from substrakt.utils.a1 import a1_in_range

            if region.sheet == evidence.locator.sheet and a1_in_range(
                evidence.locator.a1.split(":")[0], region.a1
            ):
                return True
        elif isinstance(region, (PdfRect, ImageRect)) and isinstance(
            evidence.locator, (PdfRect, ImageRect)
        ):
            if region.page == evidence.locator.page and _rect_overlaps(
                evidence.locator, region, 0.3
            ):
                return True
    return False


def _type_matches(value: str, json_type: str) -> bool:
    from substrakt.extraction import normalization

    if not value:
        return False
    if json_type in ("number", "integer"):
        return normalization.amount(value) is not None
    if json_type == "date":
        return normalization.date(value) is not None
    if json_type == "boolean":
        return value.strip().lower() in ("true", "false", "ja", "nein", "wahr", "falsch")
    return True


def confidence(
    evidence: Evidence | None,
    value: str,
    json_type: str = "string",
    regions: list[Locator] | None = None,
    extractor_confidence: float = 1.0,
) -> float:
    if evidence is None:
        return 0.0

    similarity = evidence.similarity / 100.0
    ocr = evidence.ocr_confidence
    ocr_factor = ocr if ocr is not None else 1.0  
    type_factor = 1.0 if _type_matches(value, json_type) else 0.5

    
    region_factor = 1.0
    if regions and not _in_region(evidence, regions):
        region_factor = 0.5

    result = similarity * ocr_factor * type_factor * extractor_confidence * region_factor

    return round(min(1.0, max(0.0, result)), 4)


def is_grounded(confidence_value: float, threshold: float = 0.5) -> bool:
    return confidence_value >= threshold
