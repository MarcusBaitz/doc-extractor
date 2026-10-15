"""Spreadsheet and delimited-text parser."""

from __future__ import annotations

import csv
import logging
from pathlib import Path

from charset_normalizer import from_bytes
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from substrakt.errors import DocumentCorrupt, DocumentEmpty, DocumentPasswordProtected
from substrakt.model import Document, DocumentMeta, Page, SheetRange, Word

log = logging.getLogger(__name__)


# It is substantially faster but does not preserve formatting.
LARGE_FROM_CELLS = 200_000




MAX_ROWS = 100_000
MAX_COLUMNS = 1_024


def _a1(row: int, column: int) -> str:
    """1-basierte Zeile/Spalte → A1-Notation."""
    return f"{get_column_letter(column)}{row}"


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "WAHR" if value else "FALSCH"
    return str(value).strip()




# unplausibel — dort gewinnt CP1252.
PLAUSIBLE_ENCODINGS = frozenset(
    {"utf_8", "utf8", "cp1252", "windows_1252", "latin_1", "iso8859_1",
     "iso8859_15", "latin_9", "cp850", "cp437"}
)


def _detect_encoding(raw_bytes: bytes) -> str:
    import codecs

    if raw_bytes.startswith(codecs.BOM_UTF8):
        return "utf-8-sig"
    if raw_bytes.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return "utf-16"
    if raw_bytes.startswith((codecs.BOM_UTF32_LE, codecs.BOM_UTF32_BE)):
        return "utf-32"

    
    # praktisch sicher UTF-8.
    try:
        raw_bytes.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        pass

    match = from_bytes(raw_bytes).best()
    if match and match.encoding:
        normalized = match.encoding.lower().replace("-", "_")
        if normalized in PLAUSIBLE_ENCODINGS:
            return match.encoding
        log.info(
            "Rejected implausible encoding suggestion %s; using cp1252",
            match.encoding,
        )

    return "cp1252"


def _read_csv(path: Path) -> tuple[list[list[str]], str, str]:
    try:
        raw_bytes = path.read_bytes()
    except OSError as exc:
        raise DocumentCorrupt(f"CSV cannot be read: {exc}") from exc

    encoding = _detect_encoding(raw_bytes)

    try:
        raw_text = raw_bytes.decode(encoding, errors="replace")
    except LookupError as exc:
        raise DocumentCorrupt(f"Unbekannte Kodierung {encoding}: {exc}") from exc

    if not raw_text.strip():
        raise DocumentEmpty("The CSV file is empty.")

    probe = raw_text[:8192]
    try:
        dialekt = csv.Sniffer().sniff(probe, delimiters=";,\t|")
        separator = dialekt.delimiter
    except csv.Error:
        
        erste_zeile = probe.splitlines()[0] if probe.splitlines() else ""
        separator = max(";,\t|", key=erste_zeile.count)
        if erste_zeile.count(separator) == 0:
            separator = ";"

    rows = [
        [cell.strip() for cell in reihe]
        for reihe in csv.reader(raw_text.splitlines(), delimiter=separator)
    ]
    return rows[:MAX_ROWS], encoding, separator


def _without_empty_border(rows: list[list[str]]) -> list[list[str]]:
    while rows and not any(rows[-1]):
        rows.pop()

    width = max((len(r) - next((i for i, z in enumerate(reversed(r)) if z), len(r))
                  for r in rows), default=0)

    return [r[:width] + [""] * (width - len(r)) for r in rows]


def _read_excel(path: Path) -> tuple[dict[str, list[list[str]]], list[str]]:
    extension = path.suffix.lower()
    warnings: list[str] = []

    def kuerze(name: str, rows: list[list[str]]) -> list[list[str]]:
        if len(rows) > MAX_ROWS:
            warnings.append(
                f"Blatt „{name}“ hat {len(rows)} rows and was truncated at "
                f"{MAX_ROWS} abgeschnitten."
            )
            rows = rows[:MAX_ROWS]
        cleaned = _without_empty_border(rows)
        if cleaned and len(cleaned[0]) > MAX_COLUMNS:
            warnings.append(
                f"Blatt „{name}“ hat {len(cleaned[0])} columns and was truncated at "
                f"{MAX_COLUMNS} abgeschnitten."
            )
            cleaned = [r[:MAX_COLUMNS] for r in cleaned]
        return cleaned

    
    gross = path.stat().st_size > 20 * 1024 * 1024
    if extension == ".xls" or gross:
        try:
            from python_calamine import CalamineWorkbook

            mappe = CalamineWorkbook.from_path(str(path))
            return {
                name: kuerze(
                    name,
                    [
                        [_cell_to_text(z) for z in reihe]
                        for reihe in mappe.get_sheet_by_name(name).to_python()
                    ],
                )
                for name in mappe.sheet_names
            }, warnings
        except Exception as exc:
            if extension == ".xls":
                raise DocumentCorrupt(f"Excel file cannot be read: {exc}") from exc
            log.warning("calamine failed; using openpyxl: %s", exc)

    try:
        
        
        mappe = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        message = str(exc).lower()
        if "password" in message or "encrypt" in message:
            raise DocumentPasswordProtected(
                "The workbook is password-protected."
            ) from exc
        raise DocumentCorrupt(f"Excel file cannot be read: {exc}") from exc

    try:
        result: dict[str, list[list[str]]] = {}
        for sheet in mappe.worksheets:
            
            
            rows = [
                [_cell_to_text(z) for z in reihe]
                for reihe in sheet.iter_rows(max_row=MAX_ROWS + 1, values_only=True)
            ]
            result[sheet.title] = kuerze(sheet.title, rows)
        return result, warnings
    finally:
        mappe.close()


def parse(path: Path, doc_id: str, sha256: str, size_bytes: int) -> Document:
    encoding: str | None = None
    separator: str | None = None
    warnings: list[str] = []

    if path.suffix.lower() in (".csv", ".tsv"):
        rows, encoding, separator = _read_csv(path)
        sheets = {"Daten": _without_empty_border(rows)}
    else:
        sheets, warnings = _read_excel(path)

    if not sheets:
        raise DocumentEmpty("The file contains no sheets.")

    pages: list[Page] = []
    words: list[Word] = []

    for index, (name, rows) in enumerate(sheets.items()):
        max_spalten = max((len(r) for r in rows), default=0)

        for zeilen_nr, reihe in enumerate(rows, start=1):
            for spalten_nr, cell in enumerate(reihe, start=1):
                if not cell:
                    continue
                words.append(
                    Word(
                        text=cell,
                        page=index,
                        locator=SheetRange(sheet=name, a1=_a1(zeilen_nr, spalten_nr)),
                        confidence=None,  
                    )
                )

        pages.append(
            Page(
                number=index,
                
                
                # ein leeres Blatt trotzdem darstellbar bleiben muss.
                width=float(max(max_spalten, 1)),
                height=float(max(len(rows), 1)),
                rotation=0,
                sheet_name=name,
                has_text_layer=True,  
            )
        )

    
    

    if not words:
        warnings.append("The spreadsheet contains no populated cells.")

    return Document(
        id=doc_id,
        filename=path.name,
        sha256=sha256,
        kind="sheet",
        pages=pages,
        words=words,
        meta=DocumentMeta(
            size_bytes=size_bytes,
            page_count=len(pages),
            word_count=len(words),
            text_source="cells" if words else "none",
            encoding=encoding,
            delimiter=separator,
            warnings=warnings,
        ),
    )


def read_grid(path: Path, sheet_index: int) -> list[list[str]]:
    if path.suffix.lower() in (".csv", ".tsv"):
        rows, _, _ = _read_csv(path)
        return _without_empty_border(rows) if sheet_index == 0 else []

    sheets, _warnungen = _read_excel(path)
    names = list(sheets)
    if sheet_index < 0 or sheet_index >= len(names):
        return []
    return sheets[names[sheet_index]]
